from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# 프레젠테이션 생성
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# 색상 정의
COLOR_BLUE = RGBColor(31, 78, 121)
COLOR_LIGHT_BLUE = RGBColor(79, 129, 189)
COLOR_GREEN = RGBColor(155, 187, 89)
COLOR_DARK_GRAY = RGBColor(89, 89, 89)
COLOR_LIGHT_GRAY = RGBColor(242, 242, 242)

def add_title_slide(title, subtitle):
    """제목 슬라이드 추가"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 슬라이드
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BLUE

    # 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(54)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # 부제목
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = subtitle
    subtitle_frame.paragraphs[0].font.size = Pt(24)
    subtitle_frame.paragraphs[0].font.color.rgb = RGBColor(200, 200, 200)
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_content_slide(title, content_list):
    """내용 슬라이드 추가"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(255, 255, 255)

    # 헤더 (파란색)
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_BLUE
    header.line.color.rgb = COLOR_BLUE

    # 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(40)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # 내용
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True

    for i, item in enumerate(content_list):
        if i > 0:
            text_frame.add_paragraph()
        p = text_frame.paragraphs[i]
        p.text = item
        p.font.size = Pt(18)
        p.font.color.rgb = COLOR_DARK_GRAY
        p.space_before = Pt(8)
        p.space_after = Pt(8)
        p.level = 0

def add_code_slide(title, code_text):
    """코드 슬라이드 추가"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 헤더
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_BLUE
    header.line.color.rgb = COLOR_BLUE

    # 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(40)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # 코드 박스
    code_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.7))
    text_frame = code_box.text_frame
    text_frame.word_wrap = True
    p = text_frame.paragraphs[0]
    p.text = code_text
    p.font.name = 'Courier New'
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(0, 0, 0)

    # 코드 박스 배경
    code_shape = slide.shapes.add_shape(1, Inches(0.4), Inches(1.2), Inches(9.2), Inches(5.9))
    code_shape.fill.solid()
    code_shape.fill.fore_color.rgb = COLOR_LIGHT_GRAY
    code_shape.line.color.rgb = RGBColor(200, 200, 200)
    slide.shapes._spTree.remove(code_shape._element)
    slide.shapes._spTree.insert(2, code_shape._element)

# ========== 슬라이드 추가 ==========

# 1. 제목 슬라이드
add_title_slide("홈앤쇼핑 교육 현황 대시보드", "Streamlit + Supabase 환경 구성 가이드")

# 2. 목차
add_content_slide("📋 목차", [
    "✓ 프로젝트 개요",
    "✓ 사전 요구사항",
    "✓ Supabase 설정",
    "✓ 데이터 마이그레이션",
    "✓ Streamlit 설정",
    "✓ 앱 실행 및 기능",
    "✓ 트러블슈팅"
])

# 3. 프로젝트 개요
add_content_slide("📌 프로젝트 개요", [
    "목표: CSV 데이터를 Supabase DB에 저장하고 Streamlit으로 시각화",
    "",
    "기술 스택:",
    "  • Frontend: Streamlit",
    "  • Database: Supabase (PostgreSQL)",
    "  • API: Supabase REST API",
    "  • 시각화: Plotly"
])

# 4. 사전 요구사항
add_content_slide("🔧 사전 요구사항", [
    "Python 3.10 이상",
    "",
    "필수 패키지 설치:",
    "  pip install streamlit pandas plotly requests",
    "",
    "Supabase 계정 필요",
    "  (https://supabase.com)"
])

# 5. Supabase 설정 - 1부
add_content_slide("☁️ Supabase 설정 (1/2)", [
    "1. Supabase 프로젝트 생성",
    "   • https://supabase.com 방문",
    "   • 새 프로젝트 생성",
    "",
    "2. Project 정보 확인",
    "   • Project URL: https://xxx.supabase.co",
    "   • Project Ref: xxx",
    "",
    "3. API 키 획득",
    "   • Settings → API → Project API Keys",
    "   • anon 키 또는 publishable key 복사"
])

# 6. Supabase 설정 - 2부 (테이블 생성)
add_code_slide("☁️ Supabase 설정 - 테이블 생성", """CREATE TABLE ethics_training (
  id SERIAL PRIMARY KEY,
  month VARCHAR(10) NOT NULL,
  department VARCHAR(50) NOT NULL,
  course_name VARCHAR(100) NOT NULL,
  total_employees INTEGER NOT NULL,
  completed INTEGER NOT NULL,
  not_completed INTEGER NOT NULL,
  enrollment_rate DECIMAL(5, 2),
  completion_rate DECIMAL(5, 2),
  created_at TIMESTAMP DEFAULT NOW()
);""")

# 7. 데이터 마이그레이션
add_content_slide("📊 데이터 마이그레이션", [
    "CSV 파일의 데이터를 SQL로 변환",
    "",
    "변환 규칙:",
    "  • 수강진도율, 완료율: % 제거 후 숫자만 저장",
    "  • month: '1월', '2월' 등 문자 그대로 저장",
    "",
    "방법:",
    "  • Supabase MCP 도구로 INSERT 실행",
    "  • 또는 Supabase 콘솔에서 직접 입력"
])

# 8. Streamlit 설정 - 디렉토리 구조
add_content_slide("🎨 Streamlit 설정", [
    "디렉토리 구조:",
    "  project_root/",
    "  ├── app.py",
    "  ├── .streamlit/",
    "  │   └── secrets.toml",
    "  ├── ethics_training_data.csv",
    "  └── requirements.txt",
    "",
    "핵심: .streamlit/secrets.toml에 Supabase 정보 저장"
])

# 9. Streamlit Secrets 파일
add_code_slide("🔐 .streamlit/secrets.toml 작성", """[supabase]
url = "https://lxctrbmrjeypvctkpckl.supabase.co"
key = "sb_publishable_TE-_PjBSvT..."

⚠️ 주의: .gitignore에 추가하여 GitHub에 업로드 X""")

# 10. Streamlit 코드 예제
add_code_slide("💻 Streamlit에서 데이터 로드", """import streamlit as st
import requests

@st.cache_data
def load_data():
    url = st.secrets["supabase"]["url"]
    key = st.secrets["supabase"]["key"]

    headers = {'apikey': key}
    response = requests.get(
        f'{url}/rest/v1/ethics_training',
        headers=headers
    )
    return response.json()""")

# 11. 앱 실행
add_content_slide("▶️ 앱 실행", [
    "1. 패키지 설치",
    "   pip install -r requirements.txt",
    "",
    "2. Streamlit 실행",
    "   streamlit run app.py",
    "",
    "3. 브라우저 접근",
    "   http://localhost:8501",
    "",
    "✓ 자동으로 브라우저가 열림"
])

# 12. 대시보드 기능
add_content_slide("📈 대시보드 기능", [
    "✓ KPI 카드: 대상 인원, 이수자, 미이수자, 완료율",
    "",
    "✓ 데이터 테이블: 월별, 직급별 교육 현황",
    "",
    "✓ 파이 차트: 직급별 이수 현황 (3개)",
    "",
    "✓ 라인 차트: 월별 평균 완료율 추이",
    "",
    "✓ 필터: 월, 직급으로 동적 필터링"
])

# 13. 데이터 흐름
add_content_slide("🔄 데이터 흐름", [
    "CSV 파일",
    "   ↓",
    "Supabase 데이터베이스 (ethics_training)",
    "   ↓",
    "Streamlit 앱 (REST API 호출)",
    "   ↓",
    "브라우저에서 시각화",
    "",
    "폴백: Supabase 연결 실패 → CSV에서 자동 로드"
])

# 14. 트러블슈팅 - 1부
add_content_slide("🔍 트러블슈팅 (1/2)", [
    "❌ ModuleNotFoundError",
    "   → pip install -r requirements.txt",
    "",
    "❌ Supabase 연결 오류",
    "   → secrets.toml 파일 확인",
    "   → URL과 Key 정확성 확인",
    "",
    "❌ 404 Not Found (REST API)",
    "   → ethics_training 테이블 존재 확인"
])

# 15. 트러블슈팅 - 2부
add_content_slide("🔍 트러블슈팅 (2/2)", [
    "❌ 포트 8501 이미 사용 중",
    "   → streamlit run app.py --server.port 8502",
    "",
    "❌ 데이터가 표시되지 않음",
    "   → Supabase 콘솔에서 테이블 데이터 확인",
    "   → 브라우저 캐시 삭제 시도",
    "",
    "폴백 동작: CSV 파일이 자동으로 로드됨"
])

# 16. 배포 가이드
add_content_slide("🚀 배포 시 주의사항", [
    "1. .gitignore 확인",
    "   .streamlit/secrets.toml",
    "   .env 등 보안 파일 제외",
    "",
    "2. 배포 플랫폼에서 환경 변수 설정",
    "   • Streamlit Cloud",
    "   • Heroku",
    "   • Railway 등",
    "",
    "3. requirements.txt 최신화",
    "   pip freeze > requirements.txt"
])

# 17. 참고 문서
add_content_slide("📚 참고 문서", [
    "Streamlit 공식 문서",
    "  https://docs.streamlit.io/",
    "",
    "Supabase 공식 문서",
    "  https://supabase.com/docs",
    "",
    "Plotly 시각화",
    "  https://plotly.com/python/",
    "",
    "GitHub 저장소",
    "  https://github.com/mitybomb-cloud/cloud"
])

# 18. 마무리
add_title_slide("모든 설정이 완료되었습니다! 🎉", "이제 대시보드를 만나보세요!")

# 파일 저장
prs.save(r'C:\jth\setup_guide.pptx')
print("PowerPoint file created successfully: setup_guide.pptx")
